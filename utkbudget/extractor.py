"""Extract labelled budget numbers from a UTK proposal-budget spreadsheet.

The parser targets the central UTK "Proposal Budget" workbook layout
(``TCE 001 (Rev 03.09.26)`` and compatible revisions).  All of the cell
locations live in :data:`LAYOUT`-style constants near the top of this module so
that future spreadsheet revisions only require editing the row numbers here.

The public entry point is :func:`extract_budget`, which returns a
:class:`Budget` -- an ordered collection of :class:`Field` records.  Each field
carries a machine name (used to build LaTeX ``\\newcommand`` macros), the raw
value, and a ``kind`` that controls formatting and whether the field may be
summed when several budgets are merged.
"""

from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass, field
from typing import Any, Iterable, Optional

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Worksheet / cell layout
# ---------------------------------------------------------------------------

SHEET_NAME = "UTK Budget"

# Period (budget year) columns.  L-P are the five funding periods, Q is the
# row total.  The human-readable suffixes are appended to macro names so that a
# value such as the salaries total for period 2 becomes ``...YearTwo``.
PERIOD_COLS = ["L", "M", "N", "O", "P"]
PERIOD_WORDS = ["One", "Two", "Three", "Four", "Five"]
TOTAL_COL = "Q"

# Section A -- Senior Personnel
SENIOR_ROWS = range(11, 23)          # 12 senior personnel slots
SENIOR_SUBTOTAL_ROW = 23

# Section B -- Other Personnel
POSTDOC_ROWS = range(26, 29)         # 3
OTHER_PROF_ROWS = range(29, 32)      # 3
GRA_ROWS = range(32, 36)             # 4
UNDERGRAD_ROW = 36
ADMIN_ROW = 37
OTHER_STAFF_ROWS = range(38, 40)     # 2
OTHER_SUBTOTAL_ROW = 40
WAGES_TOTAL_ROW = 41

# Section C -- Fringe Benefits (mirrors the personnel ordering above)
SENIOR_FRINGE_ROWS = range(44, 56)   # 12
POSTDOC_FRINGE_ROWS = range(56, 59)  # 3
OTHER_PROF_FRINGE_ROWS = range(59, 62)  # 3
GRA_FRINGE_ROWS = range(62, 66)      # 4
UNDERGRAD_FRINGE_ROW = 66
ADMIN_FRINGE_ROW = 67
OTHER_STAFF_FRINGE_ROWS = range(68, 70)  # 2
FRINGE_TOTAL_ROW = 70
SALARY_BENEFITS_TOTAL_ROW = 71

# Section D -- Equipment
EQUIPMENT_ROWS = range(74, 79)       # 5 itemised lines
EQUIPMENT_TOTAL_ROW = 79

# Section E -- Travel
DOMESTIC_TRAVEL_ROW = 81
FOREIGN_TRAVEL_ROW = 82
TRAVEL_TOTAL_ROW = 83

# Section F -- Participant Support
PARTICIPANT_SUPPORT_ROW = 85

# Section G -- Other Direct Costs
OTHER_DIRECT_ROWS = OrderedDict([
    ("Supplies", 88),
    ("Publication", 89),
    ("Shipping", 90),
    ("Maintenance", 91),
    ("Consultant", 92),
    ("ComputerServices", 93),
    ("Subcontracts", 94),
    ("Contractual", 95),
    ("UserFacility", 96),
    ("OtherDirectLine", 97),
])
TUITION_ROW = 99
DIFFERENTIAL_TUITION_ROW = 100
MANDATORY_FEES_ROW = 101
TUITION_SUBTOTAL_ROW = 102
OTHER_DIRECT_TOTAL_ROW = 103

# Sections H-K -- Totals & indirect costs
DIRECT_TOTAL_ROW = 105
MTDC_ROW = 108
FANDA_RATE_ROW = 109            # applied F&A rate per period (fraction)
INDIRECT_TOTAL_ROW = 110
TOTAL_ROW = 111
FANDA_RATE_TYPE_CELL = "D110"   # e.g. "Research ON-Campus"

# Metadata (label / value pairs near the top of the sheet).  ``rate`` cells are
# stored as fractions in the workbook and converted to percentages on read.
META_CELLS = OrderedDict([
    ("FundingAgency", ("D1", "text")),
    ("PINames", ("D2", "text")),
    ("ProjectTitle", ("D3", "text")),
    ("ProjectStartDate", ("D5", "date")),
    ("ProjectEndDate", ("I5", "date")),
    # Assumed escalation / raise structure (used in the budget justification).
    ("SalaryInflationUT", ("D6", "rate")),
    ("SalaryInflationJFO", ("D7", "rate")),
    ("SalaryInflationGRA", ("D8", "rate")),
    ("TuitionInflation", ("D9", "rate")),
])

# Column holding the fringe-benefit / overhead rate (a fraction such as 0.35)
RATE_COL = "F"


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

# Field "kinds" understood by the formatter / merger.
KIND_MONEY = "money"   # dollar amount; summed on merge
KIND_RATE = "rate"     # percentage (stored as a percent, e.g. 35.0)
KIND_MONTHS = "months"  # person-months (a small float)
KIND_TEXT = "text"     # free text (names, titles, ...)
KIND_DATE = "date"     # datetime

SUMMABLE_KINDS = {KIND_MONEY, KIND_MONTHS}


@dataclass
class Field:
    """A single labelled value extracted from the workbook."""

    name: str
    value: Any
    kind: str = KIND_MONEY

    @property
    def summable(self) -> bool:
        return self.kind in SUMMABLE_KINDS


@dataclass
class Budget:
    """An ordered set of :class:`Field` records for one budget."""

    source: str = ""
    fields: "OrderedDict[str, Field]" = field(default_factory=OrderedDict)

    def add(self, name: str, value: Any, kind: str = KIND_MONEY) -> None:
        self.fields[name] = Field(name, value, kind)

    def add_line(self, base: str, ws, row: int, kind: str = KIND_MONEY) -> None:
        """Add the five period columns plus the row total for ``row``."""
        for word, col in zip(PERIOD_WORDS, PERIOD_COLS):
            self.add(f"{base}Year{word}", ws[f"{col}{row}"].value, kind)
        self.add(f"{base}Total", ws[f"{TOTAL_COL}{row}"].value, kind)

    def get(self, name: str) -> Any:
        f = self.fields.get(name)
        return f.value if f is not None else None

    def items(self):
        return self.fields.items()

    def __iter__(self) -> Iterable[Field]:
        return iter(self.fields.values())


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

def _rate(value: Any) -> Optional[float]:
    """Convert a stored fraction (0.35) to a percentage (35.0)."""
    if value is None:
        return None
    try:
        return float(value) * 100.0
    except (TypeError, ValueError):
        return None


def extract_budget(path: str) -> Budget:
    """Read ``path`` and return the labelled :class:`Budget`.

    ``data_only=True`` makes openpyxl return the last values Excel cached for
    each formula, so the spreadsheet must have been opened/saved in Excel (or
    another engine that evaluates formulas) for the totals to be populated.
    """
    wb = load_workbook(filename=path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"{path!r}: expected a worksheet named {SHEET_NAME!r}; "
            f"found {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]
    b = Budget(source=path)

    # -- Metadata ---------------------------------------------------------
    for name, (cell, kind) in META_CELLS.items():
        value = ws[cell].value
        if kind == KIND_RATE:
            value = _rate(value)
        b.add(name, value, kind)

    def abc(i: int) -> str:
        """0 -> 'A', 1 -> 'B', ...  (used to label repeated personnel slots)."""
        return chr(ord("A") + i)

    # -- Section A: Senior Personnel --------------------------------------
    for i, row in enumerate(SENIOR_ROWS):
        b.add(f"Senior{abc(i)}Name", ws[f"B{row}"].value, KIND_TEXT)
        b.add(f"Senior{abc(i)}Type", ws[f"C{row}"].value, KIND_TEXT)  # UT / JFO
        b.add(f"Senior{abc(i)}BaseAnnual", ws[f"D{row}"].value, KIND_MONEY)
        b.add(f"Senior{abc(i)}ApptType", ws[f"E{row}"].value, KIND_TEXT)
        b.add(f"Senior{abc(i)}PersonMonths", ws[f"F{row}"].value, KIND_MONTHS)
        b.add_line(f"Senior{abc(i)}", ws, row, KIND_MONEY)
    b.add_line("SeniorSubtotal", ws, SENIOR_SUBTOTAL_ROW, KIND_MONEY)

    # -- Section B: Other Personnel ---------------------------------------
    def personnel_group(label: str, rows: Iterable[int]) -> None:
        for i, row in enumerate(rows):
            b.add(f"{label}{abc(i)}Name", ws[f"B{row}"].value, KIND_TEXT)
            b.add(f"{label}{abc(i)}BaseMonthly", ws[f"D{row}"].value, KIND_MONEY)
            b.add(f"{label}{abc(i)}PersonMonths", ws[f"F{row}"].value, KIND_MONTHS)
            b.add_line(f"{label}{abc(i)}", ws, row, KIND_MONEY)

    personnel_group("Postdoc", POSTDOC_ROWS)
    personnel_group("OtherProf", OTHER_PROF_ROWS)
    personnel_group("GRA", GRA_ROWS)
    personnel_group("Undergrad", [UNDERGRAD_ROW])
    personnel_group("Admin", [ADMIN_ROW])
    personnel_group("OtherStaff", OTHER_STAFF_ROWS)
    b.add_line("OtherPersonnelSubtotal", ws, OTHER_SUBTOTAL_ROW, KIND_MONEY)

    # Per-category salary subtotals (derived: the sheet only totals all "other
    # personnel" together, but the justification lists each category on its own
    # line).  Summed from the per-person rows already extracted above.
    def category_subtotal(out_base: str, members: Iterable[str]) -> None:
        members = list(members)
        for word in PERIOD_WORDS + ["Total"]:
            suffix = f"Year{word}" if word != "Total" else "Total"
            total = 0.0
            for m in members:
                v = b.get(f"{m}{suffix}")
                try:
                    total += float(v) if v not in (None, "") else 0.0
                except (TypeError, ValueError):
                    pass
            b.add(f"{out_base}{suffix}", total, KIND_MONEY)

    category_subtotal("PostdocSubtotal", (f"Postdoc{abc(i)}" for i in range(3)))
    category_subtotal("OtherProfSubtotal", (f"OtherProf{abc(i)}" for i in range(3)))
    category_subtotal("GRASubtotal", (f"GRA{abc(i)}" for i in range(4)))
    category_subtotal("UndergradSubtotal", ["UndergradA"])
    category_subtotal("AdminSubtotal", ["AdminA"])
    category_subtotal("OtherStaffSubtotal", (f"OtherStaff{abc(i)}" for i in range(2)))
    # NB: base names never end in "Total"; add_line() appends the suffix, so the
    # row total below becomes the macro \<prefix>WagesTotal (not WagesTotalTotal).
    b.add_line("Wages", ws, WAGES_TOTAL_ROW, KIND_MONEY)

    # -- Section C: Fringe Benefits ---------------------------------------
    def fringe_group(label: str, rows: Iterable[int]) -> None:
        for i, row in enumerate(rows):
            b.add(f"{label}{abc(i)}FringeRate", _rate(ws[f"{RATE_COL}{row}"].value), KIND_RATE)
            b.add_line(f"{label}{abc(i)}Fringe", ws, row, KIND_MONEY)

    fringe_group("Senior", SENIOR_FRINGE_ROWS)
    fringe_group("Postdoc", POSTDOC_FRINGE_ROWS)
    fringe_group("OtherProf", OTHER_PROF_FRINGE_ROWS)
    fringe_group("GRA", GRA_FRINGE_ROWS)
    fringe_group("Undergrad", [UNDERGRAD_FRINGE_ROW])
    fringe_group("Admin", [ADMIN_FRINGE_ROW])
    fringe_group("OtherStaff", OTHER_STAFF_FRINGE_ROWS)
    b.add_line("Fringe", ws, FRINGE_TOTAL_ROW, KIND_MONEY)
    b.add_line("SalaryAndBenefits", ws, SALARY_BENEFITS_TOTAL_ROW, KIND_MONEY)

    # -- Section D: Equipment ---------------------------------------------
    for i, row in enumerate(EQUIPMENT_ROWS):
        b.add(f"Equipment{abc(i)}Name", ws[f"B{row}"].value, KIND_TEXT)
        b.add_line(f"EquipmentItem{abc(i)}", ws, row, KIND_MONEY)
    b.add_line("Equipment", ws, EQUIPMENT_TOTAL_ROW, KIND_MONEY)

    # -- Section E: Travel ------------------------------------------------
    b.add_line("DomesticTravel", ws, DOMESTIC_TRAVEL_ROW, KIND_MONEY)
    b.add_line("ForeignTravel", ws, FOREIGN_TRAVEL_ROW, KIND_MONEY)
    b.add_line("Travel", ws, TRAVEL_TOTAL_ROW, KIND_MONEY)

    # -- Section F: Participant Support -----------------------------------
    b.add_line("ParticipantSupport", ws, PARTICIPANT_SUPPORT_ROW, KIND_MONEY)

    # -- Section G: Other Direct Costs ------------------------------------
    for base, row in OTHER_DIRECT_ROWS.items():
        b.add_line(base, ws, row, KIND_MONEY)
    b.add_line("Tuition", ws, TUITION_ROW, KIND_MONEY)
    b.add_line("DifferentialTuition", ws, DIFFERENTIAL_TUITION_ROW, KIND_MONEY)
    b.add_line("MandatoryFees", ws, MANDATORY_FEES_ROW, KIND_MONEY)
    b.add_line("TuitionSubtotal", ws, TUITION_SUBTOTAL_ROW, KIND_MONEY)
    b.add_line("OtherDirect", ws, OTHER_DIRECT_TOTAL_ROW, KIND_MONEY)

    # -- Sections H-K: Totals & indirect costs ----------------------------
    b.add_line("Direct", ws, DIRECT_TOTAL_ROW, KIND_MONEY)
    b.add_line("ModifiedTotalDirectCosts", ws, MTDC_ROW, KIND_MONEY)
    b.add("FandARateType", ws[FANDA_RATE_TYPE_CELL].value, KIND_TEXT)
    for word, col in zip(PERIOD_WORDS, PERIOD_COLS):
        b.add(f"OverheadRateYear{word}", _rate(ws[f"{col}{FANDA_RATE_ROW}"].value), KIND_RATE)
    b.add_line("Indirect", ws, INDIRECT_TOTAL_ROW, KIND_MONEY)
    # Grand total -> macros \<prefix>GrandYearOne ... \<prefix>GrandTotal
    b.add_line("Grand", ws, TOTAL_ROW, KIND_MONEY)

    return b
