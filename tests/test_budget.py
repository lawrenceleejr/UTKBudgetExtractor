"""Tests for utkbudget.

Run with::

    python -m unittest discover -s tests

Extractor/LaTeX tests use a tiny synthetic workbook with known values written
into the cells the extractor reads.  Workbook-merge tests copy the shipped
template (``examples/Proposal_Budget_Basic.xlsx``) so the real formulas are
present and formula-safety can be asserted for real.
"""

import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook

from utkbudget import extractor
from utkbudget.extractor import extract_budget
from utkbudget.merge import (RowOverflow, ValueConflict, merge_budgets,
                             write_merged_workbook)
from utkbudget.texdefs import format_value, tex_prefix, write_defs
from utkbudget.justification import build_document


def make_workbook(path, total=10000.0, travel_dom=1000.0, travel_for=500.0,
                  seniors=None):
    """Write a minimal but structurally valid UTK Budget workbook.

    ``seniors`` is an optional list of names written into the senior-personnel
    rows (used to exercise the concatenating merge and its row-overflow path).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = extractor.SHEET_NAME

    # metadata
    ws["D2"] = "Dr. Test PI"
    ws["D3"] = "A & B Physics"      # contains an ampersand -> tests escaping

    for i, name in enumerate(seniors or []):
        row = list(extractor.SENIOR_ROWS)[i]
        ws[f"B{row}"] = name
        ws[f"D{row}"] = 100000
        for col in extractor.PERIOD_COLS:
            ws[f"{col}{row}"] = 50000

    def line(row, base):
        tot = 0.0
        for i, col in enumerate(extractor.PERIOD_COLS):
            v = base + i
            ws[f"{col}{row}"] = v
            tot += v
        ws[f"{extractor.TOTAL_COL}{row}"] = tot
        return tot

    line(extractor.SENIOR_SUBTOTAL_ROW, 100)
    line(extractor.WAGES_TOTAL_ROW, 200)
    line(extractor.FRINGE_TOTAL_ROW, 50)
    line(extractor.SALARY_BENEFITS_TOTAL_ROW, 250)
    line(extractor.DOMESTIC_TRAVEL_ROW, travel_dom)
    line(extractor.FOREIGN_TRAVEL_ROW, travel_for)
    line(extractor.TRAVEL_TOTAL_ROW, travel_dom + travel_for)
    line(extractor.DIRECT_TOTAL_ROW, 5000)
    line(extractor.INDIRECT_TOTAL_ROW, 2000)
    ws[f"{extractor.TOTAL_COL}{extractor.TOTAL_ROW}"] = total
    for col in extractor.PERIOD_COLS:
        ws[f"{col}{extractor.TOTAL_ROW}"] = total / len(extractor.PERIOD_COLS)

    # F&A rate + type
    for col in extractor.PERIOD_COLS:
        ws[f"{col}{extractor.FANDA_RATE_ROW}"] = 0.535
    ws[extractor.FANDA_RATE_TYPE_CELL] = "Research ON-Campus"

    wb.save(path)


TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "examples", "Proposal_Budget_Basic.xlsx")


def make_real_input(path, seniors=(), domestic_airfares=(), foreign_airfares=(),
                    subcontracts=(), gras=0, gra_lines=(), postdocs=(),
                    undergrads=(), supplies=(), tuition=None):
    """Copy the shipped template (real formulas intact) and set user inputs.

    * ``seniors``   -- ``(name, base_annual, person_months)`` per senior row.
    * ``gras``      -- N GRA rows at $2500 base / 12 months (shorthand).
    * ``gra_lines`` -- ``(base_monthly, months)`` per GRA row (explicit).
    * ``postdocs``/``undergrads`` -- ``(base_monthly, months)`` per row.
    * ``domestic_airfares``/``foreign_airfares`` -- airfare per travel row
      (3 days, 1 traveler).
    * ``subcontracts`` -- ``(institution, amount)`` per row.
    * ``supplies``  -- ``(description, period1_amount)`` per row.
    * ``tuition``   -- annual per-GRA tuition cost.
    """
    shutil.copy(TEMPLATE, path)
    wb = load_workbook(path)  # keep formulas
    ws = wb[extractor.SHEET_NAME]
    PMONTHS = ["F", "G", "H", "I", "J"]

    for i, (name, base, months) in enumerate(seniors):
        r = list(extractor.SENIOR_ROWS)[i]
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = name, "UT", base, 9
        for col in PMONTHS:
            ws[f"{col}{r}"] = months

    def other_personnel(rows, C, entries):
        for i, (base, months) in enumerate(entries):
            r = list(rows)[i]
            ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = C, base, 1
            for col in PMONTHS:
                ws[f"{col}{r}"] = months

    other_personnel(extractor.POSTDOC_ROWS, "UT", postdocs)
    gra_entries = list(gra_lines) + [(2500, 12)] * gras
    other_personnel(extractor.GRA_ROWS, "GRA", gra_entries)
    other_personnel([extractor.UNDERGRAD_ROW], "UT", undergrads)

    if tuition is not None:
        ws[f"F{extractor.TUITION_ROW}"] = tuition

    def travel(first_row, airfares):
        tr = wb["TRAVEL"]
        for i, air in enumerate(airfares):
            r = first_row + i
            tr[f"D{r}"], tr[f"E{r}"], tr[f"G{r}"] = 3, 1, air  # days, travelers, airfare
    travel(4, domestic_airfares)    # Period-1 domestic rows 4-13
    travel(16, foreign_airfares)    # Period-1 foreign rows 16-20

    if subcontracts:
        sc = wb["SUBCONTRACTS"]
        for i, (inst, amount) in enumerate(subcontracts):
            r = 3 + i
            sc[f"B{r}"], sc[f"E{r}"], sc[f"K{r}"] = inst, amount, "Y"

    if supplies:
        sp = wb["SUPPLIES "]
        for i, (desc, amount) in enumerate(supplies):
            r = 3 + i
            sp[f"A{r}"], sp[f"B{r}"] = desc, amount

    wb.save(path)
    return path


class ExtractorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = os.path.join(self.tmp, "budget.xlsx")
        make_workbook(self.path, total=10000.0)

    def test_extracts_metadata_and_totals(self):
        b = extract_budget(self.path)
        self.assertEqual(b.get("PINames"), "Dr. Test PI")
        self.assertAlmostEqual(b.get("GrandTotal"), 10000.0)
        self.assertAlmostEqual(b.get("DomesticTravelTotal"), 5 * 1000.0 + (0 + 1 + 2 + 3 + 4))
        self.assertEqual(b.get("FandARateType"), "Research ON-Campus")
        # rate stored as a percentage
        self.assertAlmostEqual(b.get("OverheadRateYearOne"), 53.5)

    def test_no_doubled_total_macros(self):
        b = extract_budget(self.path)
        # The grand total uses base "Grand" so the Q macro is GrandTotal, never
        # "TotalTotal"; and aggregate bases never end in "Total".
        self.assertIn("GrandTotal", b.fields)
        self.assertIn("WagesTotal", b.fields)
        self.assertNotIn("TotalTotal", b.fields)
        self.assertNotIn("WagesTotalTotal", b.fields)


class MergeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.a = os.path.join(self.tmp, "a.xlsx")
        self.b = os.path.join(self.tmp, "b.xlsx")
        make_workbook(self.a, total=10000.0)
        make_workbook(self.b, total=25000.0)

    def test_sum_of_grand_totals(self):
        ba = extract_budget(self.a)
        bb = extract_budget(self.b)
        merged = merge_budgets([ba, bb])
        self.assertAlmostEqual(merged.get("GrandTotal"), 35000.0)

    def test_non_summable_carried_forward(self):
        merged = merge_budgets([extract_budget(self.a), extract_budget(self.b)])
        # rates are not summed (would be 107%); first value carried instead
        self.assertAlmostEqual(merged.get("OverheadRateYearOne"), 53.5)
        self.assertEqual(merged.get("FandARateType"), "Research ON-Campus")

    def test_merged_workbook_is_formula_safe(self):
        # Build two realistic inputs from the shipped template (which has the
        # real formulas) and confirm the merge never clobbers a formula.
        a = make_real_input(os.path.join(self.tmp, "ra.xlsx"),
                            seniors=[("Alice", 100000, 3)])
        b = make_real_input(os.path.join(self.tmp, "rb.xlsx"),
                            seniors=[("Bob", 120000, 2)])
        out = os.path.join(self.tmp, "rmerged.xlsx")
        write_merged_workbook([a, b], out)
        wb = load_workbook(out, data_only=False)
        ws = wb[extractor.SHEET_NAME]
        # Every one of these is a formula in the template and must stay a formula.
        for coord in ["L11", "Q11", "L23", "L44", "L70", "L71",
                      "L81", "L88", "L94", "L99", "L105", "L108", "L110", "L111"]:
            v = ws[coord].value
            self.assertTrue(isinstance(v, str) and v.startswith("="),
                            f"{coord} should still be a formula, got {v!r}")

    def test_personnel_user_inputs_concatenated(self):
        a = make_real_input(os.path.join(self.tmp, "pa.xlsx"),
                            seniors=[("Alice", 100000, 3), ("Bob", 120000, 2)])
        b = make_real_input(os.path.join(self.tmp, "pb.xlsx"),
                            seniors=[("Carol", 90000, 1)])
        out = os.path.join(self.tmp, "pmerged.xlsx")
        self.assertEqual(write_merged_workbook([a, b], out), [])
        ws = load_workbook(out)[extractor.SHEET_NAME]
        rows = list(extractor.SENIOR_ROWS)
        self.assertEqual([ws[f"B{rows[i]}"].value for i in range(3)],
                         ["Alice", "Bob", "Carol"])
        # The base-salary *input* is copied; the salary cell stays a formula.
        self.assertEqual([ws[f"D{rows[i]}"].value for i in range(3)],
                         [100000, 120000, 90000])
        self.assertTrue(str(ws[f"L{rows[0]}"].value).startswith("="))
        # Person-months are copied for ALL five periods (the period 2-5 salary
        # formulas read columns G-J), and the raise flag in column A survives.
        for col in ["F", "G", "H", "I", "J"]:
            self.assertEqual([ws[f"{col}{rows[i]}"].value for i in range(3)],
                             [3, 2, 1], f"person-months column {col}")
        self.assertEqual(ws[f"A{rows[0]}"].value, "Yes")

    def test_subcontracts_concatenated_with_flag(self):
        # Column K (Y/N) must travel with each subcontract row or the sheet's
        # own validation breaks the MTDC base.
        a = make_real_input(os.path.join(self.tmp, "sa.xlsx"),
                            subcontracts=[("ORNL", 50000)])
        b = make_real_input(os.path.join(self.tmp, "sb.xlsx"),
                            subcontracts=[("FNAL", 30000)])
        out = os.path.join(self.tmp, "smerged.xlsx")
        write_merged_workbook([a, b], out)
        sc = load_workbook(out)["SUBCONTRACTS"]
        self.assertEqual([sc["B3"].value, sc["B4"].value], ["ORNL", "FNAL"])
        self.assertEqual([sc["E3"].value, sc["E4"].value], [50000, 30000])
        self.assertEqual([sc["K3"].value, sc["K4"].value], ["Y", "Y"])

    def test_conflicting_tuition_reported(self):
        # Two inputs both budget GRAs but disagree on the per-GRA tuition cost.
        # The formulas apply one cost to every merged GRA, so the merge must
        # flag the conflict (and carry a non-zero value).
        a = make_real_input(os.path.join(self.tmp, "ca.xlsx"), gras=1, tuition=12000)
        b = make_real_input(os.path.join(self.tmp, "cb.xlsx"), gras=1, tuition=0)
        out = os.path.join(self.tmp, "cmerged.xlsx")
        issues = write_merged_workbook([a, b], out)
        conflicts = [i for i in issues if isinstance(i, ValueConflict)]
        self.assertEqual(len(conflicts), 1)
        self.assertIn("Tuition", conflicts[0].section)
        ws = load_workbook(out)[extractor.SHEET_NAME]
        self.assertEqual(ws[f"F{extractor.TUITION_ROW}"].value, 12000)

    def test_consistent_tuition_no_conflict(self):
        a = make_real_input(os.path.join(self.tmp, "na.xlsx"), gras=2, tuition=12000)
        b = make_real_input(os.path.join(self.tmp, "nb.xlsx"), gras=1, tuition=12000)
        out = os.path.join(self.tmp, "nmerged.xlsx")
        issues = write_merged_workbook([a, b], out)
        self.assertEqual([i for i in issues if isinstance(i, ValueConflict)], [])
        # 3 GRAs at the same base consolidate into ONE line (12*3 = 36 months).
        ws = load_workbook(out)[extractor.SHEET_NAME]
        gra_rows = list(extractor.GRA_ROWS)
        self.assertEqual([ws[f"D{r}"].value for r in gra_rows],
                         [2500, None, None, None])
        self.assertEqual(ws[f"F{gra_rows[0]}"].value, 36)
        self.assertEqual(ws[f"E{gra_rows[0]}"].value, 1)

    def test_gras_consolidated_beyond_capacity(self):
        # Six GRAs at one base (more than the 4 rows) across two files must
        # consolidate to a single line -- no overflow.
        a = make_real_input(os.path.join(self.tmp, "ga.xlsx"),
                            gra_lines=[(2500, 12)] * 3, tuition=12000)
        b = make_real_input(os.path.join(self.tmp, "gb.xlsx"),
                            gra_lines=[(2500, 12)] * 3, tuition=12000)
        out = os.path.join(self.tmp, "gmerged.xlsx")
        issues = write_merged_workbook([a, b], out)
        self.assertEqual([i for i in issues if isinstance(i, RowOverflow)], [])
        ws = load_workbook(out)[extractor.SHEET_NAME]
        gra_rows = list(extractor.GRA_ROWS)
        self.assertEqual(ws[f"D{gra_rows[0]}"].value, 2500)
        self.assertEqual(ws[f"F{gra_rows[0]}"].value, 72)  # 6 GRAs * 12 months
        self.assertIsNone(ws[f"D{gra_rows[1]}"].value)

    def test_gras_distinct_bases_stay_separate_and_overflow(self):
        # Five distinct GRA base salaries cannot be consolidated -> 5 lines > 4.
        a = make_real_input(os.path.join(self.tmp, "da.xlsx"),
                            gra_lines=[(2000, 12), (2500, 12), (3000, 12)])
        b = make_real_input(os.path.join(self.tmp, "db.xlsx"),
                            gra_lines=[(3500, 12), (4000, 12)])
        out = os.path.join(self.tmp, "dmerged.xlsx")
        issues = write_merged_workbook([a, b], out)
        gra_of = [i for i in issues if isinstance(i, RowOverflow)
                  and "Graduate" in i.section]
        self.assertEqual(len(gra_of), 1)
        self.assertEqual(gra_of[0].needed, 5)
        self.assertEqual(gra_of[0].capacity, 4)
        # Distinct bases occupy the four rows, largest first.
        ws = load_workbook(out)[extractor.SHEET_NAME]
        gra_rows = list(extractor.GRA_ROWS)
        self.assertEqual([ws[f"D{r}"].value for r in gra_rows],
                         [4000, 3500, 3000, 2500])

    def test_postdocs_grouped_by_base(self):
        a = make_real_input(os.path.join(self.tmp, "poa.xlsx"),
                            postdocs=[(5000, 12), (5000, 12)])
        b = make_real_input(os.path.join(self.tmp, "pob.xlsx"),
                            postdocs=[(6000, 12)])
        out = os.path.join(self.tmp, "pomerged.xlsx")
        write_merged_workbook([a, b], out)
        ws = load_workbook(out)[extractor.SHEET_NAME]
        pd = list(extractor.POSTDOC_ROWS)
        # Two distinct bases -> two lines (largest base first); $5000 has 24 mo.
        self.assertEqual(ws[f"D{pd[0]}"].value, 6000)
        self.assertEqual(ws[f"F{pd[0]}"].value, 12)
        self.assertEqual(ws[f"D{pd[1]}"].value, 5000)
        self.assertEqual(ws[f"F{pd[1]}"].value, 24)
        self.assertIsNone(ws[f"D{pd[2]}"].value)

    def test_same_base_different_type_stay_separate(self):
        # Two $5000 post-docs escalate differently if one is UT and one JFO, so
        # they must NOT be merged onto one line (years 2-5 would be wrong).
        a = make_real_input(os.path.join(self.tmp, "xa.xlsx"), postdocs=[(5000, 12)])
        b = make_real_input(os.path.join(self.tmp, "xb.xlsx"), postdocs=[(5000, 12)])
        wb = load_workbook(b)
        wb[extractor.SHEET_NAME]["C26"] = "JFO"   # different escalation rate
        wb.save(b)
        out = os.path.join(self.tmp, "xmerged.xlsx")
        write_merged_workbook([a, b], out)
        ws = load_workbook(out)[extractor.SHEET_NAME]
        pd = list(extractor.POSTDOC_ROWS)
        self.assertEqual(ws[f"D{pd[0]}"].value, 5000)
        self.assertEqual(ws[f"D{pd[1]}"].value, 5000)
        self.assertEqual(ws[f"F{pd[0]}"].value, 12)   # each keeps its own 12 mo
        self.assertEqual(ws[f"F{pd[1]}"].value, 12)
        self.assertNotEqual(ws[f"C{pd[0]}"].value, ws[f"C{pd[1]}"].value)

    def test_undergrads_collapse_to_one_line(self):
        a = make_real_input(os.path.join(self.tmp, "ua.xlsx"),
                            undergrads=[(600, 3)])
        b = make_real_input(os.path.join(self.tmp, "ub.xlsx"),
                            undergrads=[(600, 4)])
        out = os.path.join(self.tmp, "umerged.xlsx")
        write_merged_workbook([a, b], out)
        ws = load_workbook(out)[extractor.SHEET_NAME]
        r = extractor.UNDERGRAD_ROW
        # Same base -> keep base, sum months (3 + 4 = 7).
        self.assertEqual(ws[f"D{r}"].value, 600)
        self.assertEqual(ws[f"F{r}"].value, 7)

    def test_undergrads_mixed_base_collapse_normalized(self):
        # Different bases still collapse to ONE line, normalised to base=1 with
        # months = sum(base * headcount * months) so the cost is preserved.
        a = make_real_input(os.path.join(self.tmp, "uma.xlsx"),
                            undergrads=[(600, 3)])
        b = make_real_input(os.path.join(self.tmp, "umb.xlsx"),
                            undergrads=[(700, 4)])
        out = os.path.join(self.tmp, "ummerged.xlsx")
        write_merged_workbook([a, b], out)
        ws = load_workbook(out)[extractor.SHEET_NAME]
        r = extractor.UNDERGRAD_ROW
        self.assertEqual(ws[f"D{r}"].value, 1)
        self.assertEqual(ws[f"F{r}"].value, 600 * 3 + 700 * 4)  # 4600

    def test_supplies_grouped_by_description(self):
        a = make_real_input(os.path.join(self.tmp, "sua.xlsx"),
                            supplies=[("Computers", 2000), ("Chemicals", 500)])
        b = make_real_input(os.path.join(self.tmp, "sub.xlsx"),
                            supplies=[("Computers", 1500)])
        out = os.path.join(self.tmp, "sumerged.xlsx")
        write_merged_workbook([a, b], out)
        sp = load_workbook(out)["SUPPLIES "]
        # "Computers" merges to one line (2000 + 1500); "Chemicals" separate.
        self.assertEqual(sp["A3"].value, "Computers")
        self.assertEqual(sp["B3"].value, 3500)
        self.assertEqual(sp["A4"].value, "Chemicals")
        self.assertEqual(sp["B4"].value, 500)
        self.assertIsNone(sp["A5"].value)

    def test_no_formula_cell_modified_anywhere(self):
        # The rock-solid invariant: after a merge, every formula cell on every
        # sheet is byte-for-byte identical to the template's.
        a = make_real_input(os.path.join(self.tmp, "ia.xlsx"),
                            seniors=[("Alice", 100000, 3)],
                            domestic_airfares=[500],
                            subcontracts=[("ORNL", 50000)])
        b = make_real_input(os.path.join(self.tmp, "ib.xlsx"),
                            seniors=[("Bob", 120000, 2)],
                            domestic_airfares=[700])
        out = os.path.join(self.tmp, "imerged.xlsx")
        write_merged_workbook([a, b], out)
        template = load_workbook(a)  # first input is the structural template
        merged = load_workbook(out)
        checked = 0
        for sheet in template.sheetnames:
            ws_t, ws_m = template[sheet], merged[sheet]
            for row in ws_t.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        self.assertEqual(
                            ws_m[cell.coordinate].value, v,
                            f"formula clobbered at {sheet}!{cell.coordinate}")
                        checked += 1
        self.assertGreater(checked, 1000)  # sanity: we really scanned formulas

    def test_fringe_rates_untouched(self):
        # Fringe rates (column F on the fringe rows) ship with the template and
        # must not be modified by the merge.
        a = make_real_input(os.path.join(self.tmp, "fa.xlsx"),
                            seniors=[("Alice", 100000, 3)])
        b = make_real_input(os.path.join(self.tmp, "fb.xlsx"),
                            seniors=[("Bob", 120000, 2)])
        # Put a sentinel fringe rate in the template (the first input).
        wb = load_workbook(a)
        wb[extractor.SHEET_NAME]["F44"] = 0.42
        wb.save(a)
        out = os.path.join(self.tmp, "fmerged.xlsx")
        write_merged_workbook([a, b], out)
        merged = load_workbook(out)[extractor.SHEET_NAME]
        template = load_workbook(a)[extractor.SHEET_NAME]
        # The sentinel survives, and every senior fringe rate equals the
        # template's (i.e. the merge left them all untouched).
        self.assertEqual(merged["F44"].value, 0.42)
        for r in range(44, 56):
            self.assertEqual(merged[f"F{r}"].value, template[f"F{r}"].value)

    def test_travel_consolidated_per_period(self):
        # Trips across files collapse to ONE summary row per period/kind, with
        # days=1 and travelers=1 so the subtotal formula reproduces the total.
        a = make_real_input(os.path.join(self.tmp, "ta.xlsx"),
                            domestic_airfares=[500], foreign_airfares=[1500])
        b = make_real_input(os.path.join(self.tmp, "tb.xlsx"),
                            domestic_airfares=[700])
        out = os.path.join(self.tmp, "tmerged.xlsx")
        write_merged_workbook([a, b], out)
        tr = load_workbook(out)["TRAVEL"]
        # One domestic summary row (row 4): airfare = 500 + 700 (each 1 traveler).
        self.assertEqual(tr["G4"].value, 1200)
        self.assertEqual(tr["D4"].value, 1)   # days set (>0) so the total is valid
        self.assertEqual(tr["E4"].value, 1)   # travelers
        self.assertIsNone(tr["G5"].value)     # remaining trip rows cleared
        # Foreign summary row (row 16): airfare = 1500.
        self.assertEqual(tr["G16"].value, 1500)
        # The per-row total stays a formula.
        self.assertTrue(str(tr["K4"].value).startswith("="))

    def test_row_overflow_warns(self):
        cap = len(list(extractor.SENIOR_ROWS))
        a = make_real_input(os.path.join(self.tmp, "big.xlsx"),
                            seniors=[(f"PI{i}", 80000, 1) for i in range(cap)])
        b = make_real_input(os.path.join(self.tmp, "one.xlsx"),
                            seniors=[("Overflow", 80000, 1)])
        out = os.path.join(self.tmp, "omerged.xlsx")
        overflows = write_merged_workbook([a, b], out)
        senior = [o for o in overflows if o.section == "Senior Personnel"]
        self.assertEqual(len(senior), 1)
        self.assertEqual(senior[0].capacity, cap)
        self.assertEqual(senior[0].needed, cap + 1)
        self.assertEqual(senior[0].dropped, ["Overflow"])


class TexTests(unittest.TestCase):
    def test_prefix_sanitization(self):
        self.assertEqual(tex_prefix("Proposal_Budget 2.xlsx"), "ProposalBudgetTwo")
        self.assertEqual(tex_prefix("123.xlsx"), "OneTwoThree")
        self.assertEqual(tex_prefix("!!!.xlsx"), "Budget")

    def test_format_value(self):
        from utkbudget.extractor import Field, KIND_MONEY, KIND_RATE, KIND_TEXT
        # Money renders rounded to the nearest whole dollar (half rounds up).
        self.assertEqual(format_value(Field("x", 1234.5, KIND_MONEY)), "1,235")
        self.assertEqual(format_value(Field("x", 1234.49, KIND_MONEY)), "1,234")
        self.assertEqual(format_value(Field("x", None, KIND_MONEY)), "0")
        self.assertEqual(format_value(Field("x", 53.5, KIND_RATE)), "53.5")
        self.assertEqual(format_value(Field("x", "A & B", KIND_TEXT)), r"A \& B")

    def test_round_dollar(self):
        from utkbudget.extractor import round_dollar
        self.assertEqual(round_dollar(10.49), 10)
        self.assertEqual(round_dollar(10.50), 11)
        self.assertEqual(round_dollar(10.0), 10)
        self.assertEqual(round_dollar(0), 0)

    def test_defs_and_document_macros_match(self):
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "budget.xlsx")
        make_workbook(path)
        b = extract_budget(path)
        defs = os.path.join(tmp, "b.tex")
        write_defs(b, "Test", defs)
        import re
        with open(defs) as fh:
            defs_text = fh.read()
        defined = set(re.findall(r"\\newcommand\{\\([A-Za-z]+)\}", defs_text))
        doc = build_document("T", ["b.tex"],
                             [("Test", b, "Test Budget", False, True)])
        body = re.sub(r"(?<!\\)%.*", "", doc)      # drop LaTeX comments first
        body = re.sub(r"\\input\{[^}]+\}", "", body)
        used = set(re.findall(r"\\([A-Za-z]+)\{\}", body))
        undefined = used - defined
        self.assertEqual(undefined, set(), f"undefined macros: {undefined}")

    def test_underscores_escaped_in_titles_and_prose(self):
        from utkbudget.justification import render_section
        # A section heading with an underscore must be escaped, or LaTeX won't
        # compile.
        section = render_section("Test", heading="PI_Smith_2025")
        self.assertIn(r"\section{PI\_Smith\_2025}", section)
        self.assertNotIn("PI_Smith", section)   # no raw underscore survives
        # Document title is escaped too.
        doc = build_document("Grant_Budget_A", [], [])
        self.assertIn(r"\title{Grant\_Budget\_A}", doc)


class JustificationContentTests(unittest.TestCase):
    def _budget(self):
        from utkbudget.extractor import Budget, KIND_MONEY, KIND_RATE, KIND_TEXT, KIND_MONTHS
        b = Budget(source="x")

        def money(name, v):
            b.add(name, v, KIND_MONEY)
        # Active periods 1-3 only (years 4-5 zero).
        for w, v in zip(["One", "Two", "Three", "Four", "Five"], [100, 100, 100, 0, 0]):
            money(f"GrandYear{w}", v)
        money("GrandTotal", 300)
        money("SeniorSubtotalTotal", 1000)
        money("GRASubtotalTotal", 500)
        money("PostdocSubtotalTotal", 0)      # omitted from the table
        money("FringeTotal", 400)
        money("EquipmentTotal", 0)            # -> "no equipment"
        money("ParticipantSupportTotal", 0)   # -> "none requested"
        money("TravelTotal", 60)
        money("OtherDirectTotal", 200)
        money("DirectTotal", 2000)
        money("IndirectTotal", 800)
        for w in ["One", "Two", "Three", "Four", "Five"]:
            for base in ["DomesticTravel", "ForeignTravel", "Travel"]:
                money(f"{base}Year{w}", 10)
        b.add("OverheadRateYearOne", 53.5, KIND_RATE)
        b.add("FandARateType", "Research ON-Campus", KIND_TEXT)
        for name, v in [("SalaryInflationUT", 3.0), ("SalaryInflationJFO", 4.3),
                        ("SalaryInflationGRA", 5.0), ("TuitionInflation", 2.0),
                        ("SeniorAFringeRate", 34.6), ("GRAAFringeRate", 10.8)]:
            b.add(name, v, KIND_RATE)
        b.add("SeniorAName", "Dr. Example", KIND_TEXT)
        b.add("SeniorAType", "JFO", KIND_TEXT)
        b.add("SeniorAApptType", 9, KIND_TEXT)
        money("SeniorABaseAnnual", 120000)
        money("SeniorATotal", 1000)
        for w, v in zip(["One", "Two", "Three", "Four", "Five"], [2, 2, 0, 0, 0]):
            b.add(f"SeniorAMonthsYear{w}", v, KIND_MONTHS)
        return b

    def test_section_content(self):
        from utkbudget.justification import render_section
        s = render_section("X", self._budget(), heading=None, is_sum=False, detailed=True)
        # Summary table: nonzero categories only; zero rows omitted.
        self.assertIn("Senior Personnel &", s)
        self.assertIn("Graduate Research Assistants &", s)
        self.assertNotIn("Post-docs &", s)
        self.assertNotIn("Equipment &", s)
        self.assertNotIn("Participant Support &", s)
        # F&A rate named in the row label.
        self.assertIn(r"Indirect Costs (F\&A, \XOverheadRateYearOne{}\%)", s)
        # Escalation table shows JFO (a JFO senior exists) + GRA + tuition rows.
        self.assertIn("Jointly-appointed faculty (JFO)", s)
        self.assertIn("Graduate tuition", s)
        # Senior list: only nonzero-month years (1 and 2), not 3-5.
        self.assertIn("in year 1", s)
        self.assertIn("in year 2", s)
        self.assertNotIn("in year 3", s)
        # Equipment / participant zero-value text.
        self.assertIn("No equipment is requested", s)
        self.assertIn("No participant support costs are requested", s)
        # F&A: fixed rate, no "first period".
        self.assertIn("fixed for all periods", s)
        self.assertNotIn("first period", s)
        # Travel table: only active periods (1-3), not 4-5.
        self.assertIn("Period 3", s)
        self.assertNotIn("Period 4", s)

    def test_subtitle_rendered_and_escaped(self):
        doc = build_document("Budget Justification", [], [], subtitle="Dr. Jane_Doe")
        self.assertIn(r"{\large Dr. Jane\_Doe}", doc)

    def test_faculty_summary_table(self):
        from utkbudget.justification import build_faculty_summary
        tex = build_faculty_summary([("Dr. A", 100.0, 50.0, 150.0),
                                     ("R_D Lab", 200.0, 100.0, 300.0)])
        # One row per faculty, every figure a macro reference (values live in the
        # defs file), with the label still escaped for LaTeX.
        self.assertIn(r"Dr. A & \${}\FacultySummaryDrADirect{} & "
                      r"\${}\FacultySummaryDrAIndirect{} & "
                      r"\${}\FacultySummaryDrATotal{} \\", tex)
        self.assertIn(r"R\_D Lab & \${}\FacultySummaryRDLabDirect{}", tex)
        self.assertIn(r"\textbf{Total} & \textbf{\${}\FacultySummaryGrandDirect{}} & "
                      r"\textbf{\${}\FacultySummaryGrandIndirect{}} & "
                      r"\textbf{\${}\FacultySummaryGrandTotal{}} \\", tex)
        # No hard-coded dollar amounts anywhere in the table.
        self.assertNotIn(r"\$100", tex)
        self.assertNotIn(r"\$450", tex)
        # includable via the same guard as the justifications
        self.assertIn(r"\ifdefined\budgetjustificationincluded", tex)
        self.assertIn(r"\input{\budgetjustificationpath faculty_summary_defs}", tex)

    def test_faculty_summary_totals_consistent_after_rounding(self):
        from utkbudget.justification import build_summary_defs
        # Rows are rounded once at ingestion, so the stored total equals the sum
        # of the stored rows ($10 + $10 = $20), not round(10.4 + 10.4) = 21.
        defs = build_summary_defs(by_faculty=[(None, [("A", 10.4, 0.0, 10.4),
                                                     ("B", 10.4, 0.0, 10.4)])])
        self.assertIn(r"\newcommand{\FacultySummaryADirect}{10}", defs)
        self.assertIn(r"\newcommand{\FacultySummaryGrandDirect}{20}", defs)
        self.assertNotIn("{21}", defs)

    def test_include_guard_both_modes(self):
        """Every generated .tex must emit a full document standalone and a bare
        body when included -- so simulate TeX's \\ifdefined/\\def and check."""
        import re
        from utkbudget.justification import (build_document, build_pdf_driver,
                                             build_faculty_summary,
                                             build_faculty_summary_by_year)
        GUARD = r"\budgetjustificationincluded"

        def expand(text, defined):
            text = re.sub(r"(?<!\\)%.*", "", text)   # comments are not tokens
            out = []
            tok = re.compile(r"\\ifdefined(\\[A-Za-z]+)|\\else|\\fi|"
                             r"\\(?:def|providecommand)\{?(\\[A-Za-z]+)\}?\{\}")

            def parse(i, emit):
                while i < len(text):
                    m = tok.search(text, i)
                    if not m:
                        if emit:
                            out.append(text[i:])
                        return len(text), None
                    if emit:
                        out.append(text[i:m.start()])
                    i, kind = m.end(), m.group(0)
                    if kind.startswith(r"\ifdefined"):
                        cond = m.group(1) in defined
                        i, closed = parse(i, emit and cond)
                        if closed == "else":
                            i, _ = parse(i, emit and not cond)
                    elif kind == r"\else":
                        return i, "else"
                    elif kind == r"\fi":
                        return i, "fi"
                    elif emit:
                        defined.add(m.group(2))
                return i, None

            parse(0, True)
            return "".join(out)

        b = self._budget()
        files = {
            "justification": build_document(
                "Budget Justification", ["defs.tex"],
                [("X", b, None, False, True)], subtitle="Dr. A"),
            "driver": build_pdf_driver(
                ["a_justification.tex", "b_justification.tex"],
                summaries=[("Budget Request by Faculty", "faculty_summary.tex")]),
            "summary": build_faculty_summary([("Dr. A", 1.0, 2.0, 3.0)]),
            "by_year": build_faculty_summary_by_year([("Dr. A", [1.0, 2.0])]),
        }
        for name, text in files.items():
            alone = expand(text, set())
            for k in (r"\documentclass", r"\begin{document}", r"\end{document}"):
                self.assertEqual(alone.count(k), 1, f"{name} standalone: {k}")
            included = expand(text, {GUARD})
            for k in (r"\documentclass", r"\begin{document}", r"\end{document}",
                      r"\usepackage"):
                self.assertEqual(included.count(k), 0,
                                 f"{name} included: {k} must not be emitted")
        # Standalone, the driver must define the guard before pulling children in.
        drv = re.sub(r"(?<!\\)%.*", "", files["driver"])
        self.assertLess(drv.index(r"\def" + GUARD + "{}"), drv.index(r"\input{"))

    def test_faculty_summary_by_year(self):
        from utkbudget.justification import build_faculty_summary_by_year
        tex = build_faculty_summary_by_year([
            ("Dr. A", [100.0, 110.0, 120.0, 0, 0]),
            ("Dr. B", [200.0, 210.0, 220.0, 0, 0]),
        ])
        # One column per funded year; unfunded years 4-5 are not printed.
        self.assertIn(r"Faculty & Year 1 & Year 2 & Year 3 & Total Requested \\", tex)
        self.assertNotIn("Year 4 ", tex)
        self.assertIn(r"\begin{tabular}{lrrrr}", tex)
        # Figures are macro references, not literals.
        self.assertIn(r"Dr. A & \${}\FacultySummaryDrAYearOne{} & "
                      r"\${}\FacultySummaryDrAYearTwo{} & "
                      r"\${}\FacultySummaryDrAYearThree{} & "
                      r"\${}\FacultySummaryDrAYearsTotal{} \\", tex)
        self.assertIn(r"\textbf{Total} & \textbf{\${}\FacultySummaryGrandYearOne{}}",
                      tex)
        self.assertIn(r"\ifdefined\budgetjustificationincluded", tex)
        self.assertIn(r"\input{\budgetjustificationpath faculty_summary_defs}", tex)

    def test_faculty_summary_by_year_grouped(self):
        from utkbudget.justification import build_faculty_summary_by_year
        tex = build_faculty_summary_by_year(groups=[
            ("Energy Frontier", [("Dr. A", [100.0, 110.0]),
                                 ("Dr. B", [200.0, 210.0])]),
            ("Theory Frontier", [("Dr. C", [50.0, 55.0])]),
        ])
        self.assertIn(r"\multicolumn{4}{l}{\textbf{Energy Frontier}} \\", tex)
        self.assertIn(r"\quad Dr. A & \${}\FacultySummaryEnergyFrontierDrAYearOne{}",
                      tex)
        self.assertIn(r"\textit{Energy Frontier subtotal} & "
                      r"\textit{\${}\FacultySummaryEnergyFrontierSubtotalYearOne{}}",
                      tex)

    def test_faculty_summary_grouped_by_thrust(self):
        from utkbudget.justification import build_faculty_summary
        tex = build_faculty_summary(groups=[
            ("Energy Frontier", [("Dr. A", 100.0, 50.0, 150.0),
                                 ("Dr. B", 10.0, 5.0, 15.0)]),
            ("Theory_Frontier", [("Dr. C", 200.0, 100.0, 300.0)]),
        ])
        # group headings (escaped) with the PIs indented beneath them
        self.assertIn(r"\multicolumn{4}{l}{\textbf{Energy Frontier}}", tex)
        self.assertIn(r"\multicolumn{4}{l}{\textbf{Theory\_Frontier}}", tex)
        self.assertIn(r"\quad Dr. A & \${}\FacultySummaryEnergyFrontierDrADirect{}",
                      tex)
        self.assertIn(r"\textit{Energy Frontier subtotal} & "
                      r"\textit{\${}\FacultySummaryEnergyFrontierSubtotalDirect{}}",
                      tex)
        self.assertIn(r"\textbf{Total} & \textbf{\${}\FacultySummaryGrandDirect{}}",
                      tex)

    def test_summary_defs_hold_the_sums(self):
        from utkbudget.justification import build_summary_defs
        defs = build_summary_defs(
            by_faculty=[("Energy Frontier", [("Dr. A", 100.0, 50.0, 150.0),
                                             ("Dr. B", 10.0, 5.0, 15.0)]),
                        ("Theory Frontier", [("Dr. C", 200.0, 100.0, 300.0)])],
            by_year=[("Energy Frontier", [("Dr. A", [60.0, 90.0]),
                                          ("Dr. B", [5.0, 10.0])]),
                     ("Theory Frontier", [("Dr. C", [100.0, 200.0])])])
        # Per-PI figures.
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierDrADirect}{100}",
                      defs)
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierDrATotal}{150}", defs)
        # Per-thrust subtotals: 100+10 direct, 50+5 indirect, 150+15 total.
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierSubtotalDirect}{110}",
                      defs)
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierSubtotalTotal}{165}",
                      defs)
        # Grand totals across every thrust.
        self.assertIn(r"\newcommand{\FacultySummaryGrandDirect}{310}", defs)
        self.assertIn(r"\newcommand{\FacultySummaryGrandTotal}{465}", defs)
        # By-year figures share the same per-PI stem, with year suffixes.
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierDrAYearOne}{60}", defs)
        self.assertIn(r"\newcommand{\FacultySummaryEnergyFrontierDrAYearsTotal}{150}",
                      defs)
        self.assertIn(r"\newcommand{\FacultySummaryGrandYearOne}{165}", defs)
        self.assertIn(r"\newcommand{\FacultySummaryGrandYearsTotal}{465}", defs)
        # Load-once guard, so a document including both tables (each \input-ing
        # this file) does not redefine every macro.
        self.assertIn(r"\ifdefined\facultysummarydefsloaded\else", defs)
        self.assertIn(r"\def\facultysummarydefsloaded{}", defs)
        self.assertTrue(defs.rstrip().endswith(r"\fi"))

    def test_summary_macro_names_keep_surnames_and_dedupe(self):
        from utkbudget.justification import build_summary_defs
        # "Dr. Lee" must not lose its surname to filename-extension stripping,
        # and two people who sanitise to the same stem must not collide.
        defs = build_summary_defs(by_faculty=[
            (None, [("Dr. Lee", 1.0, 0.0, 1.0), ("Dr. Lee", 2.0, 0.0, 2.0)])])
        self.assertIn(r"\newcommand{\FacultySummaryDrLeeDirect}{1}", defs)
        self.assertIn(r"\newcommand{\FacultySummaryDrLeeTwoDirect}{2}", defs)


class CliJustificationTests(unittest.TestCase):
    def test_justifications_input_shared_defs_and_driver(self):
        from utkbudget.cli import run
        tmp = tempfile.mkdtemp()
        indir = os.path.join(tmp, "in")
        os.makedirs(indir)
        # PI names deliberately differ from the filenames.
        make_real_input(os.path.join(indir, "PI_Smith.xlsx"),
                        seniors=[("Alice", 100000, 2)])
        make_real_input(os.path.join(indir, "PI_Jones.xlsx"),
                        seniors=[("Bob", 90000, 1)])
        outdir = os.path.join(tmp, "out")
        run(indir, outdir)

        smith = os.path.join(outdir, "PI_Smith_justification.tex")
        defs = os.path.join(outdir, "PI_Smith.tex")
        combined = os.path.join(outdir, "justification.tex")
        driver = os.path.join(outdir, "all_justifications.tex")
        for p in (smith, defs, combined, driver):
            self.assertTrue(os.path.exists(p), f"missing {p}")

        with open(smith) as fh:
            smith_text = fh.read()
        # It \input{}s its dedicated defs file (not inlined) and uses that file's
        # unique macro prefix; the include guard lets it be dropped into a proposal.
        self.assertIn(r"\input{\budgetjustificationpath PI_Smith}", smith_text)
        self.assertNotIn(r"\newcommand{", smith_text)  # defs live in the defs file
        self.assertIn(r"\ifdefined\budgetjustificationincluded", smith_text)
        self.assertIn(r"\PISmith", smith_text)          # unique-prefix macro
        # Rendered content: generic title + PI subtitle, no other PI.
        self.assertIn("Budget Justification", smith_text)
        self.assertIn("Alice", smith_text)
        self.assertNotIn("Bob", smith_text)

        # The dedicated defs file holds the uniquely-prefixed definitions.
        with open(defs) as fh:
            self.assertIn(r"\newcommand{\PISmithGrandTotal}", fh.read())

        # The driver pulls in every justification, and is itself includable: it
        # only opens/closes the document (and defines the guard for its children)
        # when it is NOT being included, tracked by its own standalone marker.
        with open(driver) as fh:
            driver_text = fh.read()
        self.assertIn(r"\ifdefined\budgetjustificationincluded\else", driver_text)
        self.assertIn(r"\def\budgetjustificationincluded{}", driver_text)
        self.assertIn(r"\def\budgetjustificationstandalone{}", driver_text)
        self.assertIn("\\ifdefined\\budgetjustificationstandalone\n"
                      "\\end{document}", driver_text)
        self.assertIn(r"\input{\budgetjustificationpath PI_Smith_justification}",
                      driver_text)
        self.assertIn(r"\input{\budgetjustificationpath PI_Jones_justification}",
                      driver_text)
        # No \usepackage in the body -- illegal once included in a parent.
        self.assertNotIn(r"\usepackage{import}", driver_text)
        # Reading order: both summary tables at the very top, then the combined
        # justification, then the individual ones.
        order = [driver_text.index(needle) for needle in (
            r"\input{\budgetjustificationpath faculty_summary}",
            r"\input{\budgetjustificationpath faculty_summary_by_year}",
            r"\input{\budgetjustificationpath justification}",
            r"\input{\budgetjustificationpath PI_Smith_justification}")]
        self.assertEqual(order, sorted(order), "driver sections are out of order")
        # The summaries get headings here (the summary files emit none themselves).
        self.assertIn(r"\section*{Budget Request by Faculty}", driver_text)
        self.assertIn(r"\section*{Budget Request by Faculty and Year}", driver_text)

        # The combined justification carries the summed budget.
        with open(combined) as fh:
            combined_text = fh.read()
        self.assertIn("Combined", combined_text)

        # The faculty summary lists one row per PI (their final asks), with the
        # figures pulled from the shared summary defs file.
        summary = os.path.join(outdir, "faculty_summary.tex")
        self.assertTrue(os.path.exists(summary))
        with open(summary) as fh:
            summary_text = fh.read()
        self.assertIn("Alice", summary_text)
        self.assertIn("Bob", summary_text)
        self.assertIn(r"\textbf{Total}", summary_text)
        self.assertIn(r"\input{\budgetjustificationpath faculty_summary_defs}",
                      summary_text)
        self.assertIn(r"\FacultySummaryGrandTotal{}", summary_text)

        # ...and the by-year summary is written alongside it.
        by_year = os.path.join(outdir, "faculty_summary_by_year.tex")
        self.assertTrue(os.path.exists(by_year))
        with open(by_year) as fh:
            by_year_text = fh.read()
        self.assertIn("Alice", by_year_text)
        self.assertIn("Year 1", by_year_text)
        self.assertIn(r"\FacultySummaryGrandYearOne{}", by_year_text)

        # The sums themselves live in one defs file both tables \input, so the
        # numbers can be regenerated without touching either table's text.
        sdefs = os.path.join(outdir, "faculty_summary_defs.tex")
        self.assertTrue(os.path.exists(sdefs))
        with open(sdefs) as fh:
            sdefs_text = fh.read()
        for macro in (r"\newcommand{\FacultySummaryAliceDirect}",
                      r"\newcommand{\FacultySummaryAliceYearOne}",
                      r"\newcommand{\FacultySummaryGrandTotal}",
                      r"\newcommand{\FacultySummaryGrandYearsTotal}"):
            self.assertIn(macro, sdefs_text)

    def test_grouped_output_is_flat_with_unique_names(self):
        from utkbudget.cli import run
        tmp = tempfile.mkdtemp()
        indir = os.path.join(tmp, "in")
        # Two thrusts, each with a like-named spreadsheet: flattening must not
        # let one overwrite the other.
        for thrust, pi in (("Energy Frontier", "Alice"),
                           ("Theory Frontier", "Bob")):
            os.makedirs(os.path.join(indir, thrust))
            make_real_input(os.path.join(indir, thrust, "PI_Budget.xlsx"),
                            seniors=[(pi, 100000, 2)])
        outdir = os.path.join(tmp, "out")
        run(indir, outdir)

        # Every output sits directly in outdir -- no sub-directories at all.
        self.assertEqual([d for d in os.listdir(outdir)
                          if os.path.isdir(os.path.join(outdir, d))], [])
        names = sorted(os.listdir(outdir))
        # The colliding stem was disambiguated rather than overwritten.
        self.assertIn("PI_Budget.tex", names)
        self.assertIn("PI_Budget_2.tex", names)
        self.assertIn("PI_Budget_justification.tex", names)
        self.assertIn("PI_Budget_2_justification.tex", names)
        # Both PIs survive, one per justification.
        with open(os.path.join(outdir, "PI_Budget_justification.tex")) as fh:
            first = fh.read()
        with open(os.path.join(outdir, "PI_Budget_2_justification.tex")) as fh:
            second = fh.read()
        self.assertNotEqual(first, second)
        self.assertEqual({"Alice" in first, "Alice" in second}, {True, False})
        # The driver references both, by bare name (same flat directory).
        with open(os.path.join(outdir, "all_justifications.tex")) as fh:
            driver = fh.read()
        for stem in ("PI_Budget_justification", "PI_Budget_2_justification"):
            self.assertIn(r"\input{\budgetjustificationpath " + stem + "}", driver)
        self.assertNotIn("/", driver.split(r"\begin{document}")[-1])


class ProvenanceTests(unittest.TestCase):
    def test_defs_header_has_timestamp_and_commit(self):
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "budget.xlsx")
        make_workbook(path)
        defs = os.path.join(tmp, "b.tex")
        write_defs(extract_budget(path), "Test", defs)
        with open(defs) as fh:
            head = fh.read(400)
        self.assertIn("Auto-generated by UTKBudgetExtractor", head)
        self.assertIn("Generator commit:", head)

    def test_document_header_before_documentclass(self):
        doc = build_document("T", [], [])
        # Provenance comments must precede \documentclass (valid LaTeX).
        self.assertTrue(doc.lstrip().startswith("%"))
        self.assertLess(doc.index("Generator commit:"), doc.index("\\documentclass"))


if __name__ == "__main__":
    unittest.main()
